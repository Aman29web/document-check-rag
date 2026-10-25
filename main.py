from fastapi import FastAPI, UploadFile,HTTPException, File, Form
from fastapi.responses import JSONResponse, FileResponse
from pathlib import Path
from pydantic import BaseModel
import datetime
import hashlib
import uuid
import pymupdf
import io
from docx import Document
import ollama

import rag

app = FastAPI(title="DocumentChat")

# Local Ollama model used to answer questions
OLLAMA_MODEL = "granite4.1:3b"
ollama_client = ollama.AsyncClient()


# If Ollama is not running, or a model is missing, return a clear is error
# from any endpoint instead of crashing
@app.exception_handler(ConnectionError)
async def ollama_not_running(request, error):
    return JSONResponse(
        status_code=503,
        content={"detail": "Cannot reach Ollama. Is the Ollama app running?"}
    )


@app.exception_handler(ollama.ResponseError)
async def ollama_failed(request, error):
    return JSONResponse(
        status_code=502,
        content={"detail": f"Ollama error: {error.error}"}
    )


@app.get("/")
def home():
    return {"message":"Document is running!"}


# The chat web page
@app.get("/app", include_in_schema=False)
def chat_page():
    return FileResponse(Path(__file__).parent / "static" / "index.html")


@app.post("/upload-pdf")
async def upload_pdf(file: UploadFile = File(...)):

    # Check file extension
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400,
            detail="Please upload a PDF file"
        )

    # Read uploaded file
    file_bytes = await file.read()

    # Open PDF from memory
    try:
        pdf_document = pymupdf.open(stream=file_bytes, filetype="pdf")
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="This PDF file is damaged or not a real PDF"
        )

    # Extract text page by page
    extracted_text = ""

    for page in pdf_document:
        extracted_text += page.get_text() + "\n"

    page_count = pdf_document.page_count
    pdf_document.close()

    return {
        "filename": file.filename,
        "pages": page_count,
        "extracted_text": extracted_text
    }



@app.post("/upload-docx")
async def upload_docx(file: UploadFile = File(...)):

    # Check file extension
    if not file.filename.lower().endswith(".docx"):
        raise HTTPException(
            status_code=400,
            detail="Please upload a Word .docx file"
        )

    # Read uploaded file
    file_bytes = await file.read()

    # Open Word document from memory
    try:
        word_document = Document(io.BytesIO(file_bytes))
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="This Word file is damaged or not a real .docx file"
        )

    # Extract paragraphs
    extracted_text = ""

    for paragraph in word_document.paragraphs:
        if paragraph.text.strip():
            extracted_text += paragraph.text + "\n"

    # Extract tables
    for table in word_document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )
            extracted_text += row_text + "\n"

    return {
        "filename": file.filename,
        "extracted_text": extracted_text
    }


def read_pdf_pages(file_bytes):
    # Open PDF from memory
    try:
        pdf_document = pymupdf.open(stream=file_bytes, filetype="pdf")
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="This PDF file is damaged or not a real PDF"
        )

    # Extract text page by page, keeping the page number
    pages = []

    for page_number, page in enumerate(pdf_document, start=1):
        page_text = page.get_text()

        # A scanned page is just a picture, so it has no text.
        # Use OCR (Tesseract) to read the text from the picture instead.
        if not page_text.strip():
            try:
                ocr_page = page.get_textpage_ocr(language="eng", dpi=300, full=True)
                page_text = page.get_text(textpage=ocr_page)
            except RuntimeError:
                pass  # Tesseract is not installed; leave the page empty

        pages.append((page_number, page_text))

    pdf_document.close()
    return pages


def read_docx_text(file_bytes):
    # Open Word document from memory
    try:
        word_document = Document(io.BytesIO(file_bytes))
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="This Word file is damaged or not a real .docx file"
        )

    # Extract paragraphs
    extracted_text = ""

    for paragraph in word_document.paragraphs:
        if paragraph.text.strip():
            extracted_text += paragraph.text + "\n"

    # Extract tables
    for table in word_document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )
            extracted_text += row_text + "\n"

    return extracted_text


def read_document_pages(filename, file_bytes):
    # Returns a list of (page_number, text); Word files have no pages, so None
    if filename.lower().endswith(".pdf"):
        pages = read_pdf_pages(file_bytes)
    elif filename.lower().endswith(".docx"):
        pages = [(None, read_docx_text(file_bytes))]
    else:
        raise HTTPException(
            status_code=400,
            detail="Please upload a PDF or Word .docx file"
        )

    if not any(text.strip() for _, text in pages):
        raise HTTPException(
            status_code=400,
            detail="No text found in this file, even after trying OCR"
        )

    return pages


async def ask_llm(system_prompt, user_message, history=None):
    # history is a list of earlier {"question", "answer"} turns, oldest first
    messages = [{"role": "system", "content": system_prompt}]

    for turn in history or []:
        messages.append({"role": "user", "content": turn["question"]})
        messages.append({"role": "assistant", "content": turn["answer"]})

    messages.append({"role": "user", "content": user_message})

    response = await ollama_client.chat(
        model=OLLAMA_MODEL,
        messages=messages,
        options={"num_ctx": 16384, "temperature": 0}
    )
    return response.message.content


@app.post("/ask")
async def ask(file: UploadFile = File(...), question: str = Form(...)):

    # Read uploaded file
    file_bytes = await file.read()
    pages = read_document_pages(file.filename, file_bytes)

    # Join the pages, marking where each PDF page starts
    extracted_text = ""
    for page_number, page_text in pages:
        if page_number is not None:
            extracted_text += f"--- Page {page_number} ---\n"
        extracted_text += page_text + "\n"

    # Ask the local LLM, giving it only the document text to answer from
    system_prompt = (
        "You answer questions using only the document provided. "
        "If the answer is not in the document, say "
        "\"I could not find this in the document.\" "
        "If the document has page markers, mention the page number "
        "where you found the answer."
    )
    answer = await ask_llm(
        system_prompt,
        f"Document:\n{extracted_text}\n\nQuestion: {question}"
    )

    return {
        "filename": file.filename,
        "question": question,
        "answer": answer
    }


# ---------- RAG: upload once, ask many times ----------

@app.post("/documents")
async def upload_document(file: UploadFile = File(...)):

    file_bytes = await file.read()

    # If this exact file was uploaded before, reuse it instead of storing it twice
    file_hash = hashlib.sha256(file_bytes).hexdigest()
    existing_id = rag.find_document_by_hash(file_hash)
    if existing_id:
        return {
            "document_id": existing_id,
            "filename": file.filename,
            "already_uploaded": True
        }

    # Read the file, split it into chunks, and store them in ChromaDB
    pages = read_document_pages(file.filename, file_bytes)
    document_id, chunk_count = await rag.add_document(file.filename, pages, file_hash)

    return {
        "document_id": document_id,
        "filename": file.filename,
        "chunks": chunk_count,
        "already_uploaded": False
    }


@app.get("/documents")
def get_documents():
    return rag.list_documents()


@app.delete("/documents/{document_id}")
def remove_document(document_id: str):
    if not rag.delete_document(document_id):
        raise HTTPException(status_code=404, detail="Document not found")
    return {"deleted": document_id}


class Question(BaseModel):
    question: str
    document_id: str | None = None      # leave empty to search all documents
    top_k: int = 5                      # how many chunks to give the LLM
    conversation_id: str | None = None  # leave empty to start a new conversation


# Conversations are kept in memory, so they are lost when the server restarts
conversations = {}

# How many earlier question/answer turns the LLM sees
HISTORY_TURNS = 6


async def make_standalone_question(history, question):
    # A follow-up like "and on page 5?" can't be searched on its own,
    # so ask the LLM to rewrite it as a full question first
    # The conversation is given as plain text (not as chat turns) so the
    # LLM rewrites the question instead of trying to answer it
    system_prompt = (
        "You rewrite follow-up questions. You never answer them. "
        "Given a conversation and a follow-up question, rewrite the follow-up "
        "so it makes sense on its own, using names and topics from the "
        "conversation. Keep the user's own words where you can, and keep "
        "'I' and 'my' as they are (do not change them to 'you'). "
        "Reply with only the rewritten question, ending with '?'.\n"
        "Example: after a conversation about paid leave, "
        "'how many carry over?' becomes 'How many days of paid leave can carry over?'"
    )

    conversation_text = ""
    for turn in history:
        conversation_text += f"User: {turn['question']}\nAssistant: {turn['answer']}\n"

    rewritten = await ask_llm(
        system_prompt,
        f"Conversation:\n{conversation_text}\nFollow-up question: {question}\n\n"
        "Rewritten question:"
    )

    # Keep only the first line, in case the LLM adds anything after it
    return rewritten.strip().splitlines()[0]


@app.post("/documents/ask")
async def ask_documents(body: Question):

    # Continue an existing conversation, or start a new one
    if body.conversation_id:
        if body.conversation_id not in conversations:
            raise HTTPException(
                status_code=404,
                detail="Conversation not found (the server may have restarted). "
                       "Leave conversation_id empty to start a new one."
            )
        conversation_id = body.conversation_id
    else:
        conversation_id = str(uuid.uuid4())
        conversations[conversation_id] = []

    history = conversations[conversation_id][-HISTORY_TURNS:]

    # Turn follow-up questions into full questions before searching
    search_question = body.question
    if history:
        search_question = await make_standalone_question(history, body.question)

    # Find the chunks most related to the question
    matches = await rag.search(search_question, body.document_id, body.top_k)

    if not matches:
        raise HTTPException(
            status_code=404,
            detail="No documents found. Upload one with POST /documents first."
        )

    # Label each chunk so the LLM can say where the answer came from
    context = ""
    for number, match in enumerate(matches, start=1):
        source = match["filename"]
        if match["page"] is not None:
            source += f", page {match['page']}"
        context += f"[Source {number}] ({source})\n{match['text']}\n\n"

    # Quoting first makes the small model copy from the sources instead of guessing.
    # Today's date lets it understand words like "current" and "Present".
    today = datetime.date.today().strftime("%d %B %Y")
    system_prompt = (
        f"Today's date is {today}. "
        "You answer questions using only the sources provided.\n"
        "First, copy the exact sentence or line from the sources that answers the question, "
        "starting with 'Quote:'. Then give the answer, starting with 'Answer:', and cite the "
        "source, like [Source 2].\n"
        "Dates like 'June 2023 – Present' mean that is still happening today. "
        "When the question says 'I', 'me' or 'my', it means the person the "
        "document is about (for example, the person whose resume it is). "
        "If the answer is not in the sources, say \"I could not find this in the documents.\""
    )

    # The chat history is NOT sent here: the small model copies the style of
    # earlier answers and stops reading the sources carefully. The rewritten
    # question already carries what the conversation was about.
    answer = await ask_llm(
        system_prompt,
        f"Sources:\n{context}Question: {search_question}"
    )

    # Remember this turn for follow-up questions
    conversations[conversation_id].append({"question": body.question, "answer": answer})

    return {
        "conversation_id": conversation_id,
        "question": body.question,
        "searched_for": search_question,
        "answer": answer,
        "sources": matches
    }
